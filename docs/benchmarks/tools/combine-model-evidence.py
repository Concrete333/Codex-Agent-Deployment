"""Rebuild the cross-vendor Markdown reference without modifying workbooks.

Requires Python 3.10+ and openpyxl (read-only extraction).
Run from any directory; --check verifies the committed/generated document.
"""
import argparse
import hashlib
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
BENCH = ROOT / "docs" / "benchmarks"
CLAUDE = BENCH / "claude-model-evidence-2026-09-08.xlsx"
RAW = BENCH / "openai-metrics-raw-2026-09-08.json"
GPT = BENCH / "GPT-model-efficiency-2026-09-08.xlsx"
OUTPUT = ROOT / "docs" / "model-performance-comparison.md"


def records(sheet):
    headers = [c.value for c in sheet[4]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=5, values_only=True) if row[0]]


def load_rows():
    raw = json.loads(RAW.read_text(encoding="utf-8"))
    gpt = openpyxl.load_workbook(GPT, read_only=True, data_only=True)
    comparison = {r[0]: r for r in gpt["Comparison"].iter_rows(min_row=6, values_only=True)}
    rows = []
    for r in raw["rows"]:
        ref = comparison[r["short_name"]]
        assert abs(ref[1] - r["intelligence_index"]) <= 0.051
        assert ref[4] == r.get("cost_per_task_usd")
        rows.append(dict(name=r["short_name"], vendor="OpenAI", family=r["family"],
                         effort=r["reasoning_effort"], index=r["intelligence_index"],
                         status=ref[2], cost=r.get("cost_per_task_usd"),
                         output=r.get("output_tokens_per_task"), reasoning=r.get("reasoning_share_of_tokens"),
                         minutes=r.get("time_per_task_min"), input_price=r.get("price_input_per_1m"),
                         output_price=r.get("price_output_per_1m"), cache_price=r.get("price_cache_hit_per_1m"),
                         context=r.get("context_window"), omniscience=r.get("omniscience_index"),
                         evaluations={e["name"]: r.get(e["col"]) for e in raw["evalCols"]}))
    gpt.close()
    claude = openpyxl.load_workbook(CLAUDE, read_only=True, data_only=True)
    data = records(claude["Model Data"])
    benchmarks = {r["Model"]: r for r in records(claude["Benchmarks"])}
    assert len(data) == 16 and len(rows) == 22
    for r in data:
        assert r["Model"].startswith("Claude ")
        row = dict(name=r["Model"], vendor="Anthropic", family=r["Family"], effort=r["Effort"],
                   index=r["Intelligence Index"], status="Source score; no status flag", cost=r["Cost / Task ($)"],
                   output=r["Output Tokens / Task"], reasoning=r["Reasoning Share"], minutes=r["Time / Task (min)"],
                   input_price=r["Input $/M"], output_price=r["Output $/M"], cache_price=r["Cache Hit $/M"],
                   context=r["Context Window"], omniscience=r["Omniscience Index"],
                   evaluations={e["name"]: benchmarks[r["Model"]].get(e["name"]) for e in raw["evalCols"]})
        wrong = (1-r["Accuracy"]) * (1-r["Non-hallucination rate"])
        assert abs(wrong-r["Wrong answers per question"]) < 1e-9
        rows.append(row)
    claude.close()
    assert len({r["name"] for r in rows}) == 38
    for r in rows:
        acc = r["evaluations"].get("AA-Omniscience Accuracy")
        nh = r["evaluations"].get("AA-Omniscience Non-Hallucination Rate")
        r["wrong"] = (1-acc)*(1-nh) if acc is not None and nh is not None else None
        if r["wrong"] is not None and r["omniscience"] is not None:
            assert abs(100*(acc-r["wrong"])-r["omniscience"]) < 0.011, r["name"]
        r["frontier"] = "n/a"
        if r["cost"] is not None and "estimat" not in r["status"].lower():
            dominated = any(o["cost"] is not None and "estimat" not in o["status"].lower() and
                            o["cost"] <= r["cost"] and o["index"] >= r["index"] and
                            (o["cost"] < r["cost"] or o["index"] > r["index"]) for o in rows)
            r["frontier"] = "No" if dominated else "Yes"
    return sorted(rows, key=lambda r: (r["cost"] is None, r["cost"] or 0, -r["index"])), raw["evalCols"]


def num(value, places=1, percent=False):
    return "—" if value is None else f"{value * (100 if percent else 1):,.{places}f}" + ("%" if percent else "")


def table(headers, data):
    def line(row):
        return "| " + " | ".join(str(v).replace("|", "\\|") for v in row) + " |"
    return "\n".join([line(headers), line(["---"]*len(headers)), *map(line, data)])


def ranked(rows, value, descending=False):
    """Sort on the unrounded metric, missing last; exact ties use model name."""
    def key(row):
        metric = value(row)
        return (metric is None, (-metric if descending else metric) if metric is not None else 0,
                row["name"])
    return sorted(rows, key=key)


def build():
    rows, evaluations = load_rows()
    parts = ["# GPT and Claude model performance\n\nEvidence snapshot: 8 September 2026. Combined reference prepared 9 September 2026. Maintenance/research only; do not load during operational skill use.",
             "## Sources and reconciliation\n\n22 GPT configurations from the raw chart capture, checked against the corrected GPT workbook, plus 16 configurations from the Claude-only workbook. All source configuration labels, including **with fallback**, are retained. The source is [Artificial Analysis](https://artificialanalysis.ai/models), Intelligence Index v4.3."]
    for p in (RAW, GPT, CLAUDE):
        parts.append(f"- [{p.name}](benchmarks/{p.name}) — SHA-256 `{hashlib.sha256(p.read_bytes()).hexdigest()}`")
    parts.append("\nThe older `agent-model-selection-2026-09-08.xlsx` is retained as historical material, not used to generate this comparison. Its Task Fit knowledge-reliability row ranks the conditional non-hallucination rate and must not be used as a cross-vendor fabrication ranking. This document recomputes wrong answers on a common denominator and avoids treating composite dominance as a universal task recommendation.")
    parts.append("## How to interpret the numbers\n\n- Costs are API-priced benchmark task averages, not Codex/Claude allowance consumption or a quote for repository work. Do not combine subscription percentages into a dollar ranking. Listed token prices are the capture's prices, not a live price feed.\n- Output includes reasoning and answer tokens; it excludes input. Per-task averages are weighted across evaluations. Startup context, cache lifetime, reviews, retries, integration and helper calls can change real cost.\n- Intelligence Index is a composite, not coding success. Small differences do not establish a reliable winner; the source notes caution against interpreting differences around one point.\n- `—` means unavailable, never zero. Six GPT scores are marked estimated. The Claude workbook has no explicit score-status field; do not infer that every pairing was evaluated.\n- Claude Fable results explicitly include fallback in the benchmark configuration. Our wrapper does not configure fallback, so those scores are not a measured guarantee for a direct Fable-only worker.\n- A cost-frontier flag considers only captured composite score and task cost, across both vendors. It ignores task-specific strengths, uncertainty, permissions, account access and orchestration overhead.\n- Wrong answers per question = `(1 - accuracy) × (1 - non-hallucination rate)`. The published non-hallucination rate is conditional on non-correct responses. The derived measure reproduces the captured Omniscience Index via `100 × (accuracy - wrong/question)`. It measures knowledge answers, not source verification or coding defects.\n- Percentages below are benchmark scores, not probabilities of success on your assignment.")
    parts.append("## Reading the rankings\n\nEach table states its sort metric and direction. Sorts use unrounded source values, with missing measurements last and exact ties ordered by configuration name. Display rounding can hide small differences. Higher benchmark scores and lower costs are favorable on their own metric, not an overall best-to-worst model verdict. Lower token use or shorter duration does not establish equal quality. Frontier calculations also use unrounded values.")
    parts.append("## Overall capability\n\nSorted by Intelligence Index, highest first. Cost and token use are context, not secondary ranking criteria.\n\n" + table(
        ["Configuration", "Index", "Score status", "$/task", "Output tokens/task", "Reasoning share", "Combined frontier"],
        [[r["name"], num(r["index"]), r["status"], num(r["cost"],2), num(r["output"],0), num(r["reasoning"],1,True), r["frontier"]] for r in ranked(rows, lambda r: r["index"], True)]))
    for title, field, label, places, note in (
        ("Benchmark task cost", "cost", "$/task", 2, "Lowest cost first; this is not cost per accepted repository task."),
        ("Output-token use", "output", "Output tokens/task", 0, "Fewest output tokens first, including reasoning. This is token volume, not capability-adjusted efficiency."),
        ("Benchmark duration", "minutes", "Minutes/task", 2, "Shortest duration first. Elapsed time is not a routing priority."),
    ):
        parts.append(f"## {title}\n\n{note}\n\n" + table(
            ["Configuration", label, "Index", "Score status"],
            [[r["name"], num(r[field], places), num(r["index"]), r["status"]]
             for r in ranked(rows, lambda r: r[field])]))

    featured = ("Terminal-Bench v4.0", "SciCode", "AA-LCR v1.1", "AA-Omniscience Accuracy")
    parts.append("## Coding, long-context and knowledge rankings\n\nEach metric has its own descending ranking. Scores are shown as percentages with six decimal places to retain the precision of the full evaluation tables. The composite score-status flag does not establish whether each individual evaluation was measured.")
    for metric in featured:
        parts.append(f"### {metric}\n\nHighest score first.\n\n" + table(
            ["Configuration", metric, "$/task"],
            [[r["name"], num(r["evaluations"].get(metric),6,True), num(r["cost"],2)]
             for r in ranked(rows, lambda r: r["evaluations"].get(metric), True)]))
    parts.append("### Wrong answers per question\n\nLowest wrong-answer share first. Accuracy is shown alongside it: declining more questions is not the same as answering more correctly.\n\n" + table(
        ["Configuration", "Wrong/question", "Knowledge accuracy", "$/task"],
        [[r["name"], num(r["wrong"],3,True), num(r["evaluations"].get("AA-Omniscience Accuracy"),3,True), num(r["cost"],2)]
         for r in ranked(rows, lambda r: r["wrong"])]))

    parts.append("## Captured API token rates\n\nEach rate is ranked separately, lowest first, in USD per million tokens. A cheaper input rate need not mean a cheaper complete task. Context is the captured AA field, not an account/runtime entitlement. Expand a rate to see all configurations.")
    for label, field in (("Input $/M", "input_price"), ("Output $/M", "output_price"), ("Cache-hit $/M", "cache_price")):
        parts.append(f"<details>\n<summary>{label} — lowest first</summary>\n\n" + table(
            ["Configuration", label, "Context tokens"],
            [[r["name"], num(r[field],2), num(r["context"],0)] for r in ranked(rows, lambda r: r[field])]) + "\n\n</details>")

    # Keep every evaluation once, giving unrelated metrics independent rankings.
    parts.append("## Other evaluation rankings\n\nExpand a metric for its highest-to-lowest ranking. Scores retain source units (mostly fractions) to six decimal places. The four featured evaluations above are not repeated here.")
    for evaluation in evaluations:
        metric = evaluation["name"]
        if metric in featured:
            continue
        warning = ("This rate is conditional on non-correct responses. Do not read its descending order as a ranking of wrong answers per question; use the derived ranking above.\n\n"
                   if metric == "AA-Omniscience Non-Hallucination Rate" else "")
        parts.append(f"<details>\n<summary>{metric} — highest first</summary>\n\n" + warning + table(
            ["Configuration", metric, "$/task"],
            [[r["name"], num(r["evaluations"].get(metric),6), num(r["cost"],2)]
             for r in ranked(rows, lambda r: r["evaluations"].get(metric), True)]) + "\n\n</details>")
    parts.append("## Selection implications\n\nThe [interpretation notes](model-selection-analysis.md) explain how these results inform suggested starting points. The [operational model guide](../references/model-selection.md) contains choices and constraints without loading these tables. Rebuild with `python docs/benchmarks/tools/combine-model-evidence.py`; use `--check` to verify reproducibility. This process reads the workbooks without altering them.")
    return "\n\n".join(parts) + "\n"


if __name__ == "__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check",action="store_true")
    args=parser.parse_args()
    text=build()
    if args.check:
        assert OUTPUT.read_text(encoding="utf-8") == text, "Combined reference is stale"
        print("38 configurations reconciled; combined reference matches sources.")
    else:
        OUTPUT.write_text(text,encoding="utf-8")
        print(f"Wrote {OUTPUT.name}: 22 GPT + 16 Claude configurations.")
